#!/usr/bin/env python3
"""
build_index.py — regenerate the Payment Coverage explorer from the data workbook.

It reads the spreadsheet (Offerings / Countries / License tabs), validates every
row, anonymizes provider names to psp1, psp2, ... (using a stable mapping file),
and rewrites the embedded data inside index.html. The app's HTML/CSS/JS is taken
from the existing index.html, so any future design changes just work — only the
data block is replaced.

It also writes a second file (default: index.tilda.html) — a paste-ready block
for a Tilda HTML widget: no doctype/html/head/body, with all styles scoped to a
container so they can't affect the rest of the Tilda page.

USAGE
    python build_index.py DATA.xlsx [--template index.html] [--out index.html]
                                    [--mapping psp_alias.json]

    Defaults: --template index.html  --out index.html  --mapping psp_alias.json
    (So the simplest call, run in the folder containing index.html, is:
         python build_index.py Payment_coverage.xlsx )

REQUIREMENTS
    Python 3.8+  and  openpyxl   ->   pip install openpyxl

NOTES
  * Nothing is written if validation fails — bad data can't ship.
  * psp_alias.json remembers the real-name -> pspN mapping so numbering stays
    consistent across runs; new providers get the next free number. Keep this
    file; it is your private key to which psp is which. It is NOT needed by the
    published page.
  * A method whose Type is not one of the known categories is rejected, because
    the page has no color defined for it (adding a category is a code change).
"""

import argparse, json, re, sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("ERROR: openpyxl is not installed. Run:  pip install openpyxl")

KNOWN_TYPES = {"bank", "card", "ewallet", "mobile", "cash", "crypto"}
DIRECTIONS  = {"in", "out", "both"}
MODES       = {"include", "exclude", "all"}
STATUSES    = {"blocked", "caution"}

# Cosmetic method-name fixes applied on the way in (display only).
NAME_FIX = {"MobileMoney": "Mobile Money", "ApplePay": "Apple Pay", "GooglePay": "Google Pay"}

OFFERING_COLS = ["PSP", "Method", "Type", "Direction", "Mode", "Countries (ISO codes)"]


def cell(v):
    return "" if v is None else str(v).strip()


def val(row, i):
    """Safely read row[i]; blank/short rows (common in spreadsheet exports) return ''."""
    return row[i] if (i is not None and i < len(row)) else None


def scope_css(css, root):
    """Prefix every selector with `root` so the styles can't leak into a host page.
    Handles @media/@supports groups; leaves :root variables global; maps body/html
    and the universal * onto the container."""
    out = []
    i, n = 0, len(css)
    while i < n:
        j = css.find("{", i)
        if j == -1:
            break
        prelude = css[i:j].strip()
        depth, k = 1, j + 1
        while k < n and depth > 0:
            if css[k] == "{": depth += 1
            elif css[k] == "}": depth -= 1
            k += 1
        body = css[j + 1:k - 1]
        if prelude.startswith("@"):
            low = prelude.lower()
            if low.startswith("@media") or low.startswith("@supports"):
                out.append(prelude + "{" + scope_css(body, root) + "}")
            else:
                out.append(prelude + "{" + body + "}")   # @font-face, @keyframes, etc.
        else:
            scoped = []
            for s in (p.strip() for p in prelude.split(",") if p.strip()):
                if s == ":root":            scoped.append(":root")
                elif s == "*":              scoped += [root, root + " *"]
                elif s in ("html", "body"): scoped.append(root)
                else:                       scoped.append(root + " " + s)
            seen = set()
            uniq = [x for x in scoped if not (x in seen or seen.add(x))]
            out.append(", ".join(uniq) + "{" + body + "}")
        i = k
    return "".join(out)


def make_tilda(full_html, root_id="pcne"):
    """Turn a full index.html into a paste-ready fragment for a Tilda HTML block:
    no doctype/html/head/body, styles scoped to #root_id, ids namespaced to avoid
    collisions with the host page."""
    font_links = [l for l in re.findall(r"<link[^>]*>", full_html)
                  if "fonts.googleapis" in l or "gstatic" in l]
    style = re.search(r"<style>(.*?)</style>", full_html, re.S).group(1)
    body = re.search(r"<body>(.*?)</body>", full_html, re.S).group(1)
    scoped = scope_css(style, "#" + root_id)
    for _id in ("tabs", "search", "list", "detail", "toggleNames"):
        body = body.replace(f'id="{_id}"', f'id="{root_id}-{_id}"')
        body = body.replace(f'getElementById("{_id}")', f'getElementById("{root_id}-{_id}")')
    return ("\n".join(font_links) + "\n<style>\n" + scoped + "\n</style>\n"
            + f'<div id="{root_id}">\n' + body.strip() + "\n</div>\n")


def split_codes(v):
    v = cell(v)
    if not v:
        return []
    return [c.strip().upper() for c in v.replace("\n", ",").split(",") if c.strip()]


def read_sheet(wb, name):
    if name not in wb.sheetnames:
        sys.exit(f"ERROR: sheet '{name}' not found. Tabs present: {wb.sheetnames}")
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        sys.exit(f"ERROR: sheet '{name}' is empty.")
    header = [cell(h) for h in rows[0]]
    return header, rows[1:]


def build(xlsx_path, template_path, out_path, mapping_path, tilda_path=None):
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)

    # ---- Countries (reference + validation set) ----
    chead, crows = read_sheet(wb, "Countries")
    idx = {n: chead.index(n) for n in ("ISO", "Country", "Region") if n in chead}
    if len(idx) < 3:
        sys.exit(f"ERROR: Countries tab needs ISO, Country, Region columns. Found: {chead}")
    countries, valid_iso = {}, set()
    for r in crows:
        iso = cell(val(r, idx["ISO"]))
        if not iso:
            continue
        countries[iso] = {"name": cell(val(r, idx["Country"])), "region": cell(val(r, idx["Region"]))}
        valid_iso.add(iso)

    # ---- License ----
    lhead, lrows = read_sheet(wb, "License")
    lidx = {n: lhead.index(n) for n in ("ISO", "Status") if n in lhead}
    if "ISO" not in lidx:
        sys.exit(f"ERROR: License tab needs an ISO column. Found: {lhead}")
    license, errors = {}, []
    for n, r in enumerate(lrows, start=2):
        iso = cell(val(r, lidx["ISO"]))
        if not iso:
            continue
        status = cell(val(r, lidx.get("Status"))).lower() if "Status" in lidx else "blocked"
        status = status or "blocked"
        if status not in STATUSES:
            errors.append(f"License row {n}: status '{status}' not in {sorted(STATUSES)}")
        if iso not in valid_iso:
            errors.append(f"License row {n}: ISO '{iso}' not in Countries tab")
        license[iso] = {"level": ("caution" if status == "caution" else "blocked"),
                        "reason": "Restricted by our license"}

    # ---- Offerings ----
    ohead, orows = read_sheet(wb, "Offerings")
    missing = [c for c in OFFERING_COLS if c not in ohead]
    if missing:
        sys.exit(f"ERROR: Offerings tab missing columns: {missing}. Found: {ohead}")
    oi = {c: ohead.index(c) for c in OFFERING_COLS}
    order, grouped = [], {}
    for n, r in enumerate(orows, start=2):
        psp = cell(val(r, oi["PSP"]))
        method = cell(val(r, oi["Method"]))
        if not psp and not method:
            continue  # blank spacer row
        typ = cell(val(r, oi["Type"])).lower()
        direction = cell(val(r, oi["Direction"])).lower()
        mode = cell(val(r, oi["Mode"])).lower()
        codes = split_codes(val(r, oi["Countries (ISO codes)"]))

        if not psp:      errors.append(f"Offerings row {n}: missing PSP")
        if not method:   errors.append(f"Offerings row {n}: missing Method")
        if typ not in KNOWN_TYPES:
            errors.append(f"Offerings row {n} ({psp}/{method}): Type '{typ}' unknown. "
                          f"Allowed: {sorted(KNOWN_TYPES)}. (New categories need a code change.)")
        if direction not in DIRECTIONS:
            errors.append(f"Offerings row {n} ({psp}/{method}): Direction '{direction}' not in {sorted(DIRECTIONS)}")
        if mode not in MODES:
            errors.append(f"Offerings row {n} ({psp}/{method}): Mode '{mode}' not in {sorted(MODES)}")
        bad = [c for c in codes if c not in valid_iso]
        if bad:
            errors.append(f"Offerings row {n} ({psp}/{method}): unknown ISO codes {bad}")
        if mode == "all":
            if codes:
                print(f"  note: row {n} ({psp}/{method}) mode=all — ignoring the listed countries.")
            mode, codes = "exclude", []   # everywhere == exclude nothing

        if psp not in grouped:
            grouped[psp] = []
            order.append(psp)
        grouped[psp].append({"method": NAME_FIX.get(method, method),
                             "type": typ, "direction": direction, "mode": mode, "countries": codes})

    if errors:
        print("\nVALIDATION FAILED — nothing was written. Fix these and re-run:\n")
        for e in errors:
            print("  -", e)
        sys.exit(1)

    # ---- anonymize with a persistent, stable mapping ----
    mp = Path(mapping_path)
    alias = json.loads(mp.read_text()) if mp.exists() else {}
    nxt = len(alias) + 1
    for psp in order:
        if psp not in alias:
            alias[psp] = f"psp{nxt}"
            nxt += 1
    mp.write_text(json.dumps(alias, ensure_ascii=False, indent=2))
    anon = [{"name": alias[psp], "offerings": grouped[psp]} for psp in order]

    # ---- inject into the template ----
    tpl = Path(template_path)
    if not tpl.exists():
        sys.exit(f"ERROR: template '{template_path}' not found. Point --template at a copy of index.html.")
    html = tpl.read_text(encoding="utf-8")
    bundle = {"countries": countries, "psps": anon, "license": license}
    data_line = "const DATA = " + json.dumps(bundle, ensure_ascii=False, separators=(",", ":")) + ";"
    new_html, count = re.subn(r"const DATA = \{.*?\};", lambda _: data_line, html, count=1, flags=re.S)
    if count != 1:
        sys.exit("ERROR: could not find the 'const DATA = {...};' block in the template.")

    for real in order:  # safety: no real provider name may leak into the published file
        if real in new_html:
            sys.exit(f"ERROR: real provider name '{real}' would leak into output — aborting.")

    Path(out_path).write_text(new_html, encoding="utf-8")

    tilda_html = make_tilda(new_html)
    if tilda_path is None:
        p = Path(out_path)
        tilda_path = str(p.with_name(p.stem + ".tilda" + p.suffix))
    Path(tilda_path).write_text(tilda_html, encoding="utf-8")

    print("\nBuilt", out_path)
    print("      +", tilda_path, "(paste-ready block for Tilda)")
    print("  providers :", len(order), "->", ", ".join(f"{p}={alias[p]}" for p in order))
    print("  offerings :", sum(len(v) for v in grouped.values()))
    print("  countries :", len(countries), "| license blocked/caution:", len(license))
    print("  mapping   :", mapping_path)


def main():
    ap = argparse.ArgumentParser(description="Regenerate index.html from the payment-coverage workbook.")
    ap.add_argument("xlsx", help="the data workbook, e.g. Payment_coverage.xlsx")
    ap.add_argument("--template", default="index.html", help="existing index.html to use as the app shell (default: index.html)")
    ap.add_argument("--out", default="index.html", help="where to write the result (default: index.html)")
    ap.add_argument("--mapping", default="psp_alias.json", help="persistent real-name->pspN mapping (default: psp_alias.json)")
    ap.add_argument("--tilda-out", default=None, help="paste-ready Tilda fragment (default: alongside --out, e.g. index.tilda.html)")
    a = ap.parse_args()
    build(a.xlsx, a.template, a.out, a.mapping, a.tilda_out)


if __name__ == "__main__":
    main()
